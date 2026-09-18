"""
tests/test_v1_api.py — Integration and contract tests for /api/v1/ endpoints.
Verifies standardized envelope {"success": true, "data": ..., "meta": ...}
and Marshmallow schema validation errors.
"""
import pytest
from models import Student, JobRole, Skill, JobSkill, StudentSkill
from extensions import db


class TestV1HealthAndSystem:
    def test_v1_health(self, client):
        rv = client.get("/api/v1/health")
        assert rv.status_code == 200
        payload = rv.get_json()
        assert payload["success"] is True
        assert payload["data"]["status"] == "ok"

    def test_v1_cache_stats(self, client):
        rv = client.get("/api/v1/cache/stats")
        assert rv.status_code == 200
        payload = rv.get_json()
        assert payload["success"] is True
        assert "gap_cache" in payload["data"]


class TestV1Students:
    def test_create_student_success(self, client):
        rv = client.post("/api/v1/students", json={
            "name": "Stage3 Tester",
            "email": "stage3_tester@example.com",
            "course": "B.Tech CSE",
            "year": 3,
            "target_career": "Backend Developer",
        })
        assert rv.status_code == 201
        payload = rv.get_json()
        assert payload["success"] is True
        assert "id" in payload["data"]
        assert payload["data"]["name"] == "Stage3 Tester"

    def test_create_student_schema_validation_failure(self, client):
        # Invalid email and missing name
        rv = client.post("/api/v1/students", json={
            "email": "not-an-email",
        })
        assert rv.status_code == 400
        payload = rv.get_json()
        assert payload["success"] is False
        assert payload["error"] == "Validation Error"
        assert "email" in payload["messages"]
        assert "name" in payload["messages"]

    def test_list_students(self, client):
        rv = client.get("/api/v1/students")
        assert rv.status_code == 200
        payload = rv.get_json()
        assert payload["success"] is True
        assert isinstance(payload["data"], list)

    def test_get_and_update_student(self, client):
        # Create student first
        create_rv = client.post("/api/v1/students", json={
            "name": "Update Me",
            "email": "update_me@example.com",
            "course": "MCA",
            "year": 1,
        })
        sid = create_rv.get_json()["data"]["id"]

        # Get profile
        get_rv = client.get(f"/api/v1/students/{sid}")
        assert get_rv.status_code == 200
        profile = get_rv.get_json()
        assert profile["success"] is True
        assert profile["data"]["name"] == "Update Me"

        # Update profile
        update_rv = client.put(f"/api/v1/students/{sid}", json={
            "target_career": "DevOps Engineer",
            "year": 2,
        })
        assert update_rv.status_code == 200
        updated = update_rv.get_json()
        assert updated["success"] is True
        assert updated["data"]["target_career"] == "DevOps Engineer"
        assert updated["data"]["year"] == 2

    def test_add_student_skill(self, client):
        create_rv = client.post("/api/v1/students", json={
            "name": "Skill Student",
            "email": "skill_stud@example.com",
        })
        sid = create_rv.get_json()["data"]["id"]

        # Add skill
        rv = client.post(f"/api/v1/students/{sid}/skills", json={
            "skill_name": "Python",
            "proficiency": 85.0,
            "source": "self_reported",
        })
        assert rv.status_code == 200
        payload = rv.get_json()
        assert payload["success"] is True
        assert payload["data"]["proficiency"] == 85.0

    def test_add_skill_validation_range_error(self, client):
        create_rv = client.post("/api/v1/students", json={
            "name": "Range Test",
            "email": "range_test@example.com",
        })
        sid = create_rv.get_json()["data"]["id"]

        # Proficiency > 100 should fail schema validation
        rv = client.post(f"/api/v1/students/{sid}/skills", json={
            "skill_name": "Python",
            "proficiency": 150.0,
        })
        assert rv.status_code == 400
        payload = rv.get_json()
        assert payload["success"] is False
        assert "proficiency" in payload["messages"]


class TestV1SkillsAndRoles:
    def test_skills_catalog(self, client):
        rv = client.get("/api/v1/skills")
        assert rv.status_code == 200
        payload = rv.get_json()
        assert payload["success"] is True
        assert "skills" in payload["data"]
        assert "categories" in payload["data"]

    def test_roles_and_role_skills(self, client):
        rv = client.get("/api/v1/roles")
        assert rv.status_code == 200
        payload = rv.get_json()
        assert payload["success"] is True
        assert isinstance(payload["data"], list)

        if payload["data"]:
            rid = payload["data"][0]["id"]
            skills_rv = client.get(f"/api/v1/roles/{rid}/skills")
            assert skills_rv.status_code == 200
            skills_payload = skills_rv.get_json()
            assert skills_payload["success"] is True
            assert "required_skills" in skills_payload["data"]


class TestV1AssessmentsAndResumes:
    def test_assessment_flow(self, client):
        create_rv = client.post("/api/v1/students", json={
            "name": "Assess Student",
            "email": "assess_stud@example.com",
        })
        sid = create_rv.get_json()["data"]["id"]

        rv = client.post(f"/api/v1/students/{sid}/assessments", json={
            "skill_name": "Python",
            "score": 80.0,
        })
        assert rv.status_code == 201
        payload = rv.get_json()
        assert payload["success"] is True
        assert payload["data"]["proficiency"] == 80.0

    def test_projects_and_certifications(self, client):
        create_rv = client.post("/api/v1/students", json={
            "name": "Portfolio Student",
            "email": "portfolio_stud@example.com",
        })
        sid = create_rv.get_json()["data"]["id"]

        # Add project
        p_rv = client.post(f"/api/v1/students/{sid}/projects", json={
            "title": "AI Recommendation Engine",
            "description": "Flask + scikit-learn recommendation engine",
            "skills_used": ["Python", "Flask", "Machine Learning"],
        })
        assert p_rv.status_code == 201
        p_payload = p_rv.get_json()
        assert p_payload["success"] is True
        assert "id" in p_payload["data"]

        # Add certification
        c_rv = client.post(f"/api/v1/students/{sid}/certifications", json={
            "name": "AWS Certified Solutions Architect",
            "issuer": "Amazon Web Services",
        })
        assert c_rv.status_code == 201
        c_payload = c_rv.get_json()
        assert c_payload["success"] is True
        assert "id" in c_payload["data"]


class TestV1LearningAndAnalytics:
    def test_dashboard_and_analytics(self, client):
        create_rv = client.post("/api/v1/students", json={
            "name": "Analytics Student",
            "email": "analytics_stud@example.com",
        })
        sid = create_rv.get_json()["data"]["id"]

        dash_rv = client.get(f"/api/v1/students/{sid}/dashboard")
        assert dash_rv.status_code == 200
        dash_payload = dash_rv.get_json()
        assert dash_payload["success"] is True
        assert "current_skills" in dash_payload["data"]

        analytics_rv = client.get(f"/api/v1/students/{sid}/analytics")
        assert analytics_rv.status_code == 200
        analytics_payload = analytics_rv.get_json()
        assert analytics_payload["success"] is True
        assert "summary" in analytics_payload["data"]


class TestV1CoreSkillIntelligence:
    def test_gap_analysis_explainability_and_evidence(self, client):
        # Create student and role
        create_rv = client.post("/api/v1/students", json={
            "name": "Intelligence Student",
            "email": "intel_stud@example.com",
        })
        sid = create_rv.get_json()["data"]["id"]

        # Add a self-reported skill
        client.post(f"/api/v1/students/{sid}/skills", json={
            "skill_name": "Python",
            "proficiency": 30.0,
            "source": "self_reported",
        })

        # Get a populated role like Software Developer
        roles_rv = client.get("/api/v1/roles")
        roles = roles_rv.get_json()["data"]
        role = next((r for r in roles if r.get("name") == "Software Developer"), roles[0])
        rid = role["id"]

        # Run gap analysis
        analyze_rv = client.post("/api/v1/skill-gap/analyze", json={
            "student_id": sid,
            "job_role_id": rid,
        })
        assert analyze_rv.status_code == 200
        analyze_data = analyze_rv.get_json()
        assert analyze_data["success"] is True
        gaps = analyze_data["data"]["gaps"]
        assert len(gaps) > 0

        # Check Stage 5 fields
        first_gap = gaps[0]
        assert "explanation" in first_gap
        assert len(first_gap["explanation"]) > 0
        assert "evidence_summary" in first_gap
        assert "evidence_factor" in first_gap
        assert first_gap["evidence_factor"] >= 1.0

        # Verify persisted gaps endpoint includes explanation
        persisted_rv = client.get(f"/api/v1/students/{sid}/gaps?job_role_id={rid}")
        assert persisted_rv.status_code == 200
        persisted_data = persisted_rv.get_json()
        assert persisted_data["success"] is True
        assert len(persisted_data["data"]) > 0
        assert "explanation" in persisted_data["data"][0]
        assert "evidence_summary" in persisted_data["data"][0]

    def test_role_match_verification_confidence(self, client):
        create_rv = client.post("/api/v1/students", json={
            "name": "Matcher Student",
            "email": "matcher_stud@example.com",
        })
        sid = create_rv.get_json()["data"]["id"]

        # Add an assessed skill
        client.post(f"/api/v1/students/{sid}/assessments", json={
            "skill_name": "Python",
            "score": 90.0,
        })

        match_rv = client.get(f"/api/v1/students/{sid}/role-match")
        assert match_rv.status_code == 200
        match_data = match_rv.get_json()
        assert match_data["success"] is True
        matches = match_data["data"]["matches"]
        assert len(matches) > 0

        first_match = matches[0]
        assert "verification_confidence" in first_match
        assert isinstance(first_match["verification_confidence"], (int, float))
        assert "verification_breakdown" in first_match
        assert "verified" in first_match["verification_breakdown"]
        assert "explanation" in first_match
        assert len(first_match["explanation"]) > 0


class TestV1AIResumeIntelligence:
    def test_resume_upload_generates_evidence_and_audits(self, client):
        import io
        from docx import Document

        # 1. Create student
        create_rv = client.post("/api/v1/students", json={
            "name": "Resume AI Student",
            "email": "resume_ai_stud@example.com",
        })
        sid = create_rv.get_json()["data"]["id"]

        # 2. Build in-memory DOCX
        doc = Document()
        doc.add_paragraph("Skills")
        doc.add_paragraph("Python, SQL, ExoticUnmappedSpecialSkill99")
        doc.add_paragraph("Projects")
        doc.add_paragraph("Built microservice backend using Python and Docker.")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        # 3. Upload resume
        upload_rv = client.post(
            f"/api/v1/students/{sid}/resume",
            data={"file": (buf, "test_resume.docx")},
            content_type="multipart/form-data"
        )
        assert upload_rv.status_code == 201
        upload_data = upload_rv.get_json()
        assert upload_data["success"] is True
        assert "extracted_skills" in upload_data["data"]
        assert len(upload_data["data"]["extracted_skills"]) > 0
        assert "extraction_audit" in upload_data["data"]
        assert upload_data["data"]["extraction_audit"]["skills_extracted_count"] > 0
        assert upload_data["data"]["unknown_terms_count"] >= 1

        # 4. Check evidence endpoint
        evidence_rv = client.get(f"/api/v1/students/{sid}/resume/evidence")
        assert evidence_rv.status_code == 200
        ev_data = evidence_rv.get_json()["data"]
        assert len(ev_data) > 0
        first_ev = ev_data[0]
        assert "evidence_span" in first_ev
        assert "section" in first_ev
        assert first_ev["status"] == "pending"

        # 5. Check unknown skills endpoint
        unknown_rv = client.get(f"/api/v1/students/{sid}/resume/unknown-skills")
        assert unknown_rv.status_code == 200
        unk_data = unknown_rv.get_json()["data"]
        assert len(unk_data) >= 1
        assert any("exoticunmappedspecialskill99" in u["raw_term"].lower() for u in unk_data)

        # 6. Check extraction audit history
        audit_rv = client.get(f"/api/v1/students/{sid}/resume/audits")
        assert audit_rv.status_code == 200
        audits = audit_rv.get_json()["data"]
        assert len(audits) >= 1
        assert audits[0]["source_type"] == "resume"

        # 7. Verification workflow
        if len(ev_data) >= 2:
            verif_payload = {
                "verifications": [
                    {"evidence_id": ev_data[0]["id"], "action": "accept", "proficiency": 75.0},
                    {"evidence_id": ev_data[1]["id"], "action": "reject"}
                ]
            }
        else:
            verif_payload = {
                "verifications": [
                    {"evidence_id": ev_data[0]["id"], "action": "accept", "proficiency": 75.0}
                ]
            }

        verify_rv = client.post(f"/api/v1/students/{sid}/resume/verify", json=verif_payload)
        assert verify_rv.status_code == 200
        verif_res = verify_rv.get_json()["data"]
        assert verif_res["verified_count"] >= 1

        # 8. Check student skills catalog has accepted skill
        skills_rv = client.get(f"/api/v1/students/{sid}")
        stud_data = skills_rv.get_json()["data"]
        assert any(s["evidence_type"] == "resume" for s in stud_data["skills"])


class TestV1Recommendations:
    def test_recommendations_and_runs(self, client):
        import json

        create_rv = client.post("/api/v1/students", json={
            "name": "Rec Student",
            "email": "rec_stud@example.com",
        })
        sid = create_rv.get_json()["data"]["id"]

        # Add a skill and run gap analysis against Software Developer
        client.post(f"/api/v1/students/{sid}/skills", json={
            "skill_name": "Python",
            "proficiency": 30.0,
            "source": "self_reported",
        })

        roles_rv = client.get("/api/v1/roles")
        roles = roles_rv.get_json()["data"]
        role = next((r for r in roles if r.get("name") == "Software Developer"), roles[0])
        rid = role["id"]

        client.post("/api/v1/skill-gap/analyze", json={
            "student_id": sid,
            "job_role_id": rid,
        })

        # Run recommendations with default weights
        rv = client.get(f"/api/v1/students/{sid}/recommendations")
        assert rv.status_code == 200
        data = rv.get_json()
        assert data["success"] is True
        assert isinstance(data["data"], list)
        recs = data["data"]
        assert len(recs) > 0
        assert "score_breakdown" in recs[0] or "level_fit" in recs[0]

        # Run recommendations with custom weights via query parameter
        custom_weights = {
            "skill_gap_relevance": 0.50,
            "learning_level_fit": 0.10,
            "content_similarity": 0.10,
            "quality": 0.10,
            "user_preference": 0.10,
            "diversity": 0.10,
        }
        rv_custom = client.get(
            f"/api/v1/students/{sid}/recommendations?weights={json.dumps(custom_weights)}"
        )
        assert rv_custom.status_code == 200
        assert rv_custom.get_json()["success"] is True

        # Test GET /api/v1/students/<id>/recommendations/runs
        runs_rv = client.get(f"/api/v1/students/{sid}/recommendations/runs")
        assert runs_rv.status_code == 200
        runs_payload = runs_rv.get_json()
        assert runs_payload["success"] is True
        assert isinstance(runs_payload["data"], list)
        assert len(runs_payload["data"]) >= 2
        first_run = runs_payload["data"][0]
        assert "weights_used" in first_run
        assert "recommendations_snapshot" in first_run
        assert first_run["student_id"] == sid

