"""tests/test_resume_parser.py — Unit tests for resume text extraction and skill matching."""
import os
import tempfile
import pytest


class TestExtractText:
    def test_unsupported_extension_raises(self, app_ctx):
        """extract_text() should raise ValueError for unsupported extensions."""
        from services.resume_parser import extract_text
        with pytest.raises(ValueError, match="Unsupported file type"):
            extract_text("resume.txt")

    def test_nonexistent_pdf_raises(self, app_ctx):
        """extract_text() should raise an error for a non-existent PDF."""
        from services.resume_parser import extract_text
        with pytest.raises(Exception):
            extract_text("/nonexistent/path/resume.pdf")

    def test_docx_extraction(self, app_ctx):
        """extract_text() should extract text from a valid DOCX file."""
        from docx import Document
        from services.resume_parser import extract_text

        with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
            path = tmp.name

        try:
            doc = Document()
            doc.add_paragraph("Experienced Python developer with SQL and REST API expertise.")
            doc.add_paragraph("Projects: Built a Docker-based microservices application.")
            doc.save(path)

            text = extract_text(path)
            assert "Python" in text
            assert "SQL" in text
        finally:
            os.unlink(path)


class TestNormalizeText:
    def test_lowercase(self, app_ctx):
        from services.resume_parser import normalize_text
        assert normalize_text("Python SQL REST") == "python sql rest"

    def test_collapses_whitespace(self, app_ctx):
        from services.resume_parser import normalize_text
        result = normalize_text("Python   SQL\t\nREST")
        assert "  " not in result

    def test_preserves_plus_hash(self, app_ctx):
        from services.resume_parser import normalize_text
        result = normalize_text("C++ C# Python")
        assert "c++" in result
        assert "c#" in result


class TestCandidateSkills:
    def test_finds_python(self, app_ctx):
        """candidate_skills() should match 'Python' from resume text."""
        from services.resume_parser import candidate_skills
        with app_ctx.app_context():
            result = candidate_skills("Proficient in Python and SQL databases.")
            skill_names = [r["skill"] for r in result]
            assert "Python" in skill_names

    def test_confidence_in_range(self, app_ctx):
        from services.resume_parser import candidate_skills
        with app_ctx.app_context():
            result = candidate_skills("Experienced with Python, Docker, SQL.")
            for r in result:
                assert 0.0 <= r["confidence"] <= 1.0

    def test_result_fields(self, app_ctx):
        from services.resume_parser import candidate_skills
        with app_ctx.app_context():
            result = candidate_skills("Python developer with REST API experience.")
            for r in result:
                assert "skill_id" in r
                assert "skill" in r
                assert "confidence" in r
                assert "evidence" in r

    def test_no_false_positives_on_empty(self, app_ctx):
        from services.resume_parser import candidate_skills
        with app_ctx.app_context():
            result = candidate_skills("")
            assert result == []


class TestSectionExtraction:
    def test_projects_extraction(self, app_ctx):
        from services.resume_parser import extract_projects_text
        text = """
Skills
Python, SQL

Projects
E-Commerce Application using Django and REST API
Deployed a Docker container on AWS for a microservices backend

Education
B.Tech Computer Science 2022
"""
        projects = extract_projects_text(text)
        assert len(projects) >= 1
        assert any("Docker" in p or "Django" in p or "E-Commerce" in p
                   for p in projects)

    def test_certifications_extraction(self, app_ctx):
        from services.resume_parser import extract_certifications_text
        text = """
Projects
Some project

Certifications
AWS Certified Solutions Architect
Google Data Analytics Professional Certificate

Education
MCA 2024
"""
        certs = extract_certifications_text(text)
        assert len(certs) >= 1
        assert any("AWS" in c or "Google" in c for c in certs)


class TestLayeredSkillLookup:
    def test_layer1_exact(self, app_ctx):
        from services.skill_normalizer import layered_skill_lookup
        with app_ctx.app_context():
            skill, layer, conf = layered_skill_lookup("Python")
            assert skill is not None
            assert skill.name == "Python"
            assert layer == "exact"
            assert conf == 1.0

    def test_layer2_alias(self, app_ctx):
        from services.skill_normalizer import layered_skill_lookup
        with app_ctx.app_context():
            skill, layer, conf = layered_skill_lookup("py")
            assert skill is not None
            assert skill.name == "Python"
            assert layer == "alias"
            assert conf >= 0.90

    def test_layer3_rule_canonicalization(self, app_ctx):
        from services.skill_normalizer import layered_skill_lookup
        with app_ctx.app_context():
            skill, layer, conf = layered_skill_lookup("python 3.9")
            assert skill is not None
            assert skill.name == "Python"
            assert layer in ("canonical_rule", "exact", "alias")
            assert conf >= 0.85

    def test_layer4_fuzzy(self, app_ctx):
        from services.skill_normalizer import layered_skill_lookup
        with app_ctx.app_context():
            skill, layer, conf = layered_skill_lookup("Pythoon")
            assert skill is not None
            assert skill.name == "Python"
            assert layer == "fuzzy"
            assert conf >= 0.75

    def test_layer6_unknown(self, app_ctx):
        from services.skill_normalizer import layered_skill_lookup
        with app_ctx.app_context():
            skill, layer, conf = layered_skill_lookup("SuperCalifragilisticHyperDrive999")
            assert skill is None
            assert layer == "unknown"
            assert conf == 0.0


class TestEvidencedResumePipeline:
    def test_extract_skills_with_evidence_and_unknowns(self, app_ctx):
        from services.resume_parser import extract_skills_with_evidence, split_resume_into_sections
        text = """
Skills
Python, Docker, SuperRareNonExistentSkillXYZ

Projects
Built an E-Commerce backend using Python and REST APIs.
Deployed microservices on Kubernetes.

Education
B.Tech in Computer Science
"""
        with app_ctx.app_context():
            sections = split_resume_into_sections(text)
            assert "skills" in sections
            assert "projects" in sections

            mapped, unknowns = extract_skills_with_evidence(text)
            mapped_names = [m["skill"] for m in mapped]
            assert "Python" in mapped_names
            assert "Docker" in mapped_names

            # Verify evidence span and section attribution
            python_item = next(m for m in mapped if m["skill"] == "Python")
            assert "evidence_span" in python_item
            assert len(python_item["evidence_span"]) > 0
            assert "section" in python_item

            # Verify unknown skill detection
            assert len(unknowns) >= 1
            assert any("superrarenonexistentskillxyz" in u["raw_term"].lower() for u in unknowns)

